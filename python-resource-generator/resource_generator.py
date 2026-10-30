#!/usr/bin/env python3
"""Local-first AIOU study-resource pack generator; no AI API or web crawling."""
from __future__ import annotations

import argparse
import re
import sys
import urllib.request
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

ALLOWED_HOSTS = ("aiou.edu.pk",)
QUESTION_START = re.compile(
    r"^\s*(?:(?:Q(?:uestion)?\s*[-.]?\s*)?(?:\d{1,2}|[ivxlcdm]{1,6})\s*[.)\-:]|"
    r"(?:short|long|essay|multiple.choice|mcq)\s+questions?\s*[:\-])",
    re.IGNORECASE,
)
HEADING_HINTS = re.compile(
    r"^(?:unit\s+\d+|chapter\s+\d+|lesson\s+\d+|learning objectives?|"
    r"introduction|summary|conclusion|self assessment|review questions?|"
    r"objectives?|key terms|activities|exercise)\b",
    re.IGNORECASE,
)
WORD_RE = re.compile(r"[\w'-]+", re.UNICODE)


@dataclass
class SourceDocument:
    text: str
    title: str
    source_ref: str


def fetch_official_url(url: str, timeout: int = 25) -> tuple[bytes, str]:
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower().rstrip(".")
    if parsed.scheme != "https" or not any(host == h or host.endswith("." + h) for h in ALLOWED_HOSTS):
        raise ValueError("Only HTTPS URLs on official aiou.edu.pk domains are allowed.")
    request = urllib.request.Request(url, headers={"User-Agent": "AIOU-Coaching-ResourcePrototype/0.1"})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        final = urlparse(response.geturl())
        final_host = (final.hostname or "").lower().rstrip(".")
        if final.scheme != "https" or not any(final_host == h or final_host.endswith("." + h) for h in ALLOWED_HOSTS):
            raise ValueError("The source redirected outside the official AIOU domain allowlist.")
        data = response.read(35 * 1024 * 1024 + 1)
        if len(data) > 35 * 1024 * 1024:
            raise ValueError("Source file exceeds the 35 MB prototype limit.")
        return data, response.geturl()


def extract_text(path: Path | None = None, data: bytes | None = None, filename: str = "") -> str:
    suffix = (path.suffix if path else Path(filename).suffix).lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install requirements.txt to read PDF files.") from exc
        import io
        reader = PdfReader(path if path else io.BytesIO(data or b""))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            pages.append(f"\n[Page {index}]\n{page.extract_text() or ''}")
        return "\n".join(pages).strip()
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install requirements.txt to read DOCX files.") from exc
        import io
        doc = Document(path if path else io.BytesIO(data or b""))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
    if suffix in (".txt", ".md"):
        raw = path.read_bytes() if path else (data or b"")
        return raw.decode("utf-8-sig", errors="replace").strip()
    raise ValueError("Supported input types are PDF, DOCX, TXT, and MD.")


def clean_lines(text: str) -> list[str]:
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if not line or re.fullmatch(r"\[Page \d+\]", line):
            continue
        if len(line) < 3:
            continue
        lines.append(line)
    return lines


def detect_questions(lines: list[str]) -> list[str]:
    found: list[str] = []
    for line in lines:
        if QUESTION_START.match(line) and len(WORD_RE.findall(line)) >= 4:
            if line not in found:
                found.append(line)
    return found[:60]


def detect_topics(lines: list[str]) -> list[tuple[str, str]]:
    topics: list[tuple[str, str]] = []
    current = ""
    body: list[str] = []

    def flush() -> None:
        nonlocal current, body
        if current:
            topics.append((current, " ".join(body).strip()[:1400]))
        current, body = "", []

    for line in lines:
        is_heading = len(line) <= 100 and (
            HEADING_HINTS.match(line) or
            (line.isupper() and 2 <= len(WORD_RE.findall(line)) <= 12)
        )
        if is_heading:
            flush()
            current = line
        elif current and len(body) < 8:
            body.append(line)
        elif not current and len(line) <= 100 and line.istitle():
            current = line
    flush()
    seen = set()
    unique = []
    for title, summary in topics:
        key = title.casefold()
        if key not in seen and len(WORD_RE.findall(title)) >= 2:
            seen.add(key)
            unique.append((title, summary))
    return unique[:40]


def topic_prompts(topics: list[tuple[str, str]]) -> list[str]:
    prompts = []
    for title, _ in topics:
        prompts.append(f"Explain the key ideas and importance of: {title}. Use your own examples and cite your sources.")
        prompts.append(f"Prepare a short revision answer outlining the main points of: {title}.")
    return prompts[:30]


def build_outputs(source: SourceDocument, course_code: str, course_title: str, output_dir: Path) -> dict[str, Path]:
    lines = clean_lines(source.text)
    questions = detect_questions(lines)
    topics = detect_topics(lines)
    if len(WORD_RE.findall(source.text)) < 30:
        raise ValueError("Very little text was extracted. This may be a scanned PDF; OCR/review is needed.")
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_code = re.sub(r"[^A-Za-z0-9_-]+", "-", course_code).strip("-") or "course"
    stem = f"{safe_code}-study-pack"
    txt_path = output_dir / f"{stem}-extracted.txt"
    docx_path = output_dir / f"{stem}.docx"
    pdf_path = output_dir / f"{stem}.pdf"
    generated = date.today().isoformat()
    prompts = topic_prompts(topics)
    paper_questions = questions or prompts
    source_label = f"{source.title} | {source.source_ref}"

    txt_lines = [
        "AIOU Coaching — Study Resource Pack",
        f"Course: {course_title or 'Not specified'} ({course_code or 'code not specified'})",
        f"Generated: {generated}",
        f"Source: {source_label}",
        "",
        "IMPORTANT: Unofficial study aid. Verify course, semester, source and exam requirements.",
        "Use prompts to learn and write in your own words. Cite sources. These are not exam predictions.",
        "",
        "REVISION NOTES (extracted and organized; verify against the original source)",
    ]
    for title, body in topics:
        txt_lines += ["", title, body or "No adjacent paragraph was confidently detected. Review this heading in the source."]
    txt_lines += ["", "PRACTICE PAPER", "Questions below are extracted from the source where detected." if questions else
                  "No numbered questions were detected; prompts below are generated from detected headings and need human review."]
    txt_lines += [f"{i}. {q}" for i, q in enumerate(paper_questions, start=1)]
    txt_lines += ["", "SOURCE & PROVENANCE", source_label, f"Generated on: {generated}",
                  "Copyright note: access does not imply permission to republish full source text."]
    txt_path.write_text("\n".join(txt_lines) + "\n", encoding="utf-8")

    try:
        from docx import Document
        from docx.shared import Inches, Pt
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        doc.add_heading("AIOU Coaching — Study Resource Pack", 0)
        doc.add_paragraph(f"Course: {course_title or 'Not specified'} ({course_code or 'code not specified'})")
        doc.add_paragraph(f"Generated: {generated}")
        doc.add_paragraph(f"Source: {source_label}")
        note = doc.add_paragraph("Unofficial study aid. Verify against the original source. These prompts are not exam predictions.")
        note.italic = True
        doc.add_heading("Revision Notes", level=1)
        if topics:
            for title, body in topics:
                doc.add_heading(title, level=2)
                doc.add_paragraph(body or "Review this heading in the original source; no adjacent explanatory paragraph was detected.")
        else:
            doc.add_paragraph("No reliable headings detected. Use the extracted text file to review the source.")
        doc.add_heading("Practice Paper", level=1)
        if not questions:
            doc.add_paragraph("No numbered source questions were detected. Topic prompts below were generated from detected headings and need human review.")
        for i, question in enumerate(paper_questions, start=1):
            doc.add_paragraph(f"{i}. {question}")
        doc.add_heading("Source and Academic Integrity", level=1)
        doc.add_paragraph(source_label)
        doc.add_paragraph("Use these materials for study. Write assignment answers in your own words and cite sources. Do not submit extracted notes as your own original work.")
        doc.add_paragraph("Copyright: public access does not automatically permit redistribution. Check source terms before sharing.")
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10.5)
        doc.save(docx_path)
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt to export DOCX/PDF.") from exc

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from xml.sax.saxutils import escape
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="SmallNote", parent=styles["BodyText"], fontSize=8, leading=11, textColor=colors.HexColor("#555555")))
        story = [
            Paragraph("AIOU Coaching — Study Resource Pack", styles["Title"]), Spacer(1, 5*mm),
            Paragraph(f"Course: {escape(course_title or 'Not specified')} ({escape(course_code or 'code not specified')})", styles["Normal"]),
            Paragraph(f"Generated: {generated}", styles["Normal"]),
            Paragraph(f"Source: {escape(source_label)}", styles["SmallNote"]), Spacer(1, 4*mm),
            Paragraph("<b>Unofficial study aid.</b> Verify against the original source. Practice prompts are not exam predictions.", styles["SmallNote"]),
            Spacer(1, 4*mm), Paragraph("Revision Notes", styles["Heading1"])
        ]
        for title, body in topics:
            story.append(Paragraph(escape(title), styles["Heading2"]))
            story.append(Paragraph(escape(body or "Review this heading in the original source."), styles["BodyText"]))
            story.append(Spacer(1, 2*mm))
        story.append(Paragraph("Practice Paper", styles["Heading1"]))
        if not questions:
            story.append(Paragraph("No numbered source questions were detected. Prompts were generated from headings and require human review.", styles["SmallNote"]))
        for i, question in enumerate(paper_questions, start=1):
            story.append(Paragraph(f"{i}. {escape(question)}", styles["BodyText"]))
            story.append(Spacer(1, 2*mm))
        story += [Spacer(1, 4*mm), Paragraph("Source and Academic Integrity", styles["Heading1"]),
                  Paragraph(escape(source_label), styles["SmallNote"]),
                  Paragraph("Write assignment answers in your own words and cite sources. Public access does not automatically permit redistribution.", styles["SmallNote"])]
        pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        pdf.build(story)
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt to export PDF.") from exc

    return {"text": txt_path, "docx": docx_path, "pdf": pdf_path}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--input", type=Path, help="Local PDF, DOCX, TXT, or MD file")
    group.add_argument("--url", help="Explicit official AIOU HTTPS source URL (no crawling)")
    parser.add_argument("--course-code", required=True)
    parser.add_argument("--course-title", default="")
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    args = parser.parse_args(argv)
    try:
        if args.input:
            if not args.input.is_file():
                raise FileNotFoundError(f"Input file not found: {args.input}")
            source = SourceDocument(extract_text(path=args.input), args.input.name, str(args.input.resolve()))
        else:
            data, final_url = fetch_official_url(args.url)
            name = Path(urlparse(final_url).path).name or "aiou-resource"
            source = SourceDocument(extract_text(data=data, filename=name), name, final_url)
        outputs = build_outputs(source, args.course_code, args.course_title, args.output_dir)
        for kind, path in outputs.items():
            print(f"{kind.upper()}: {path}")
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
