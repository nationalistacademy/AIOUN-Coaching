import unittest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from resource_generator import clean_lines, detect_questions, detect_topics, fetch_official_url, topic_prompts


class ResourceGeneratorTests(unittest.TestCase):
    def test_clean_lines_removes_empty_and_page_markers(self):
        self.assertEqual(
            clean_lines("  Hello   world \n\n[Page 2]\n  A useful line  "),
            ["Hello world", "A useful line"],
        )

    def test_detect_questions_extracts_numbered_questions(self):
        lines = [
            "Q1. Explain the concept of business communication in detail.",
            "2) Describe the role of management in organizations.",
            "This paragraph explains management.",
        ]
        found = detect_questions(lines)
        self.assertEqual(len(found), 2)
        self.assertIn("Q1. Explain the concept of business communication in detail.", found)

    def test_topic_prompts_are_learning_tasks(self):
        prompts = topic_prompts([("Unit 1 Introduction", "A short source paragraph.")])
        self.assertTrue(prompts)
        self.assertIn("your own examples", prompts[0])

    def test_rejects_non_official_url_before_network(self):
        with self.assertRaises(ValueError):
            fetch_official_url("https://example.com/book.pdf")

    def test_detect_topics_handles_heading_and_paragraph(self):
        topics = detect_topics([
            "UNIT 1 INTRODUCTION",
            "This unit introduces the main ideas and course context.",
            "LEARNING OBJECTIVES",
            "Students will identify key terms and examples.",
        ])
        self.assertGreaterEqual(len(topics), 1)
        self.assertEqual(topics[0][0], "UNIT 1 INTRODUCTION")


    def test_end_to_end_creates_txt_docx_and_pdf_exports(self):
        from tempfile import TemporaryDirectory
        from resource_generator import SourceDocument, build_outputs
        from docx import Document
        from pypdf import PdfReader

        source = SourceDocument(
            text=(
                "UNIT 1 INTRODUCTION\\n"
                "This unit explains the foundations of business communication and its importance.\\n"
                "Q1. Explain the importance of business communication in an organization.\\n"
                "UNIT 2 MANAGEMENT\\n"
                "Management includes planning, organizing, leading and controlling resources."
            ),
            title="synthetic-ci-fixture.txt",
            source_ref="https://aiou.edu.pk/test-fixture (synthetic test content)",
        )
        with TemporaryDirectory() as tmp:
            outputs = build_outputs(source, "TEST-101", "Prototype Test Course", Path(tmp))
            self.assertEqual(set(outputs), {"text", "docx", "pdf"})
            for output in outputs.values():
                self.assertTrue(output.is_file())
                self.assertGreater(output.stat().st_size, 100)
            txt = outputs["text"].read_text(encoding="utf-8")
            self.assertIn("SOURCE & PROVENANCE", txt)
            self.assertIn("1. Q1. Explain the importance", txt)
            docx_text = "\\n".join(p.text for p in Document(outputs["docx"]).paragraphs)
            self.assertIn("Practice Paper", docx_text)
            pdf_text = "\\n".join(page.extract_text() or "" for page in PdfReader(outputs["pdf"]).pages)
            self.assertIn("Source and Academic Integrity", pdf_text)


if __name__ == "__main__":
    unittest.main()
